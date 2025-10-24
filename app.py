import streamlit as st
import pandas as pd
import json
import yaml
import io
import os
import zipfile
from datetime import datetime
from typing import Dict, List, Any, Optional, Tuple
import plotly.graph_objects as go
import base64

# Page configuration
st.set_page_config(
    page_title="Agentic Docs Builder - Flora Edition",
    page_icon="🌸",
    layout="wide",
    initial_sidebar_state="expanded"
)

# Custom CSS for Flora theme
st.markdown("""
<style>
    .main {
        background-color: #F8F8FF;
    }
    .stApp {
        background: linear-gradient(135deg, #E6E6FA 0%, #F8F8FF 100%);
    }
    .status-box {
        padding: 1rem;
        border-radius: 0.5rem;
        margin: 0.5rem 0;
        border-left: 4px solid;
    }
    .status-success {
        background-color: #D4EDDA;
        border-color: #28A745;
        color: #155724;
    }
    .status-error {
        background-color: #F8D7DA;
        border-color: #DC3545;
        color: #721C24;
    }
    .status-running {
        background-color: #FFF3CD;
        border-color: #FFC107;
        color: #856404;
    }
    .status-info {
        background-color: #D1ECF1;
        border-color: #17A2B8;
        color: #0C5460;
    }
    .metric-card {
        background: white;
        padding: 1.5rem;
        border-radius: 0.75rem;
        box-shadow: 0 2px 8px rgba(147, 112, 219, 0.1);
        border: 2px solid #E6E6FA;
    }
    .agent-card {
        background: linear-gradient(135deg, #FFFFFF 0%, #F0E6FF 100%);
        padding: 1.5rem;
        border-radius: 0.75rem;
        margin: 1rem 0;
        border: 2px solid #9370DB;
        box-shadow: 0 4px 12px rgba(147, 112, 219, 0.2);
    }
</style>
""", unsafe_allow_html=True)


# ==================== AI SERVICE ====================
class AIService:
    """Unified AI service for multiple providers"""
    
    def __init__(self, gemini_key: str = None, openai_key: str = None, grok_key: str = None):
        self.gemini_key = gemini_key or os.getenv("GEMINI_API_KEY")
        self.openai_key = openai_key or os.getenv("OPENAI_API_KEY")
        self.grok_key = grok_key or os.getenv("XAI_API_KEY")
        self._init_clients()
    
    def _init_clients(self):
        """Initialize API clients for each provider"""
        self.gemini_client = None
        self.openai_client = None
        self.grok_client = None
        
        if self.gemini_key:
            try:
                import google.generativeai as genai
                genai.configure(api_key=self.gemini_key)
                self.gemini_client = genai
            except Exception as e:
                print(f"Gemini initialization error: {e}")
        
        if self.openai_key:
            try:
                from openai import OpenAI
                self.openai_client = OpenAI(api_key=self.openai_key)
            except Exception as e:
                print(f"OpenAI initialization error: {e}")
        
        if self.grok_key:
            try:
                from xai_sdk import Client
                self.grok_client = Client(api_key=self.grok_key, timeout=3600)
            except Exception as e:
                print(f"Grok initialization error: {e}")
    
    def run_agent(self, agent: Dict[str, Any], input_text: str) -> str:
        """Run an agent with the given input"""
        provider = agent.get('provider', 'gemini')
        user_prompt = agent['user_prompt'].replace('{{input}}', input_text)
        
        if provider == 'gemini':
            return self._run_gemini(agent, user_prompt)
        elif provider == 'openai':
            return self._run_openai(agent, user_prompt)
        elif provider == 'grok':
            return self._run_grok(agent, user_prompt)
        else:
            raise ValueError(f"Unknown provider: {provider}")
    
    def _run_gemini(self, agent: Dict[str, Any], user_prompt: str) -> str:
        """Run Gemini model"""
        if not self.gemini_client:
            raise ValueError("Gemini API key not configured")
        
        model = self.gemini_client.GenerativeModel(
            model_name=agent['model'],
            generation_config={
                'temperature': agent.get('temperature', 0.7),
                'max_output_tokens': agent.get('max_tokens', 1000),
                'top_p': agent.get('top_p', 0.9),
            }
        )
        
        full_prompt = f"{agent['system_prompt']}\n\n{user_prompt}"
        response = model.generate_content(full_prompt)
        return response.text
    
    def _run_openai(self, agent: Dict[str, Any], user_prompt: str) -> str:
        """Run OpenAI model"""
        if not self.openai_client:
            raise ValueError("OpenAI API key not configured")
        
        response = self.openai_client.chat.completions.create(
            model=agent['model'],
            messages=[
                {"role": "system", "content": agent['system_prompt']},
                {"role": "user", "content": user_prompt}
            ],
            temperature=agent.get('temperature', 0.7),
            max_tokens=agent.get('max_tokens', 1000),
            top_p=agent.get('top_p', 0.9)
        )
        return response.choices[0].message.content
    
    def _run_grok(self, agent: Dict[str, Any], user_prompt: str) -> str:
        """Run Grok model using xAI SDK"""
        if not self.grok_client:
            raise ValueError("Grok API key not configured")
        
        from xai_sdk.chat import user, system
        
        chat = self.grok_client.chat.create(model=agent['model'])
        chat.append(system(agent['system_prompt']))
        chat.append(user(user_prompt))
        
        response = chat.sample(
            temperature=agent.get('temperature', 0.7),
            max_tokens=agent.get('max_tokens', 1000),
            top_p=agent.get('top_p', 0.9)
        )
        return response.content
    
    def generate_follow_up_questions(self, context: str) -> List[str]:
        """Generate follow-up questions based on the final output"""
        prompt = f"""Based on the following text, generate exactly 3 insightful follow-up questions that would help dive deeper into the topic or explore related areas:

{context}

Format your response as a simple numbered list:
1. [Question 1]
2. [Question 2]
3. [Question 3]"""
        
        try:
            if self.gemini_client:
                model = self.gemini_client.GenerativeModel('gemini-2.5-flash')
                response = model.generate_content(prompt)
                text = response.text
            elif self.openai_client:
                response = self.openai_client.chat.completions.create(
                    model="gpt-4o-mini",
                    messages=[{"role": "user", "content": prompt}],
                    temperature=0.7
                )
                text = response.choices[0].message.content
            elif self.grok_client:
                from xai_sdk.chat import user
                chat = self.grok_client.chat.create(model="grok-3-mini")
                chat.append(user(prompt))
                response = chat.sample()
                text = response.content
            else:
                return ["What are the key insights from this analysis?",
                       "How can this information be applied practically?",
                       "What additional context would be helpful?"]
            
            questions = []
            for line in text.split('\n'):
                line = line.strip()
                if line and (line[0].isdigit() or line.startswith('-') or line.startswith('•')):
                    question = line.lstrip('0123456789.-•) ').strip()
                    if question:
                        questions.append(question)
            
            return questions[:3] if len(questions) >= 3 else questions + [
                "What are the implications of these findings?",
                "How does this compare to similar cases?",
                "What would be the next steps?"
            ][:3 - len(questions)]
            
        except Exception as e:
            return [
                "What are the main takeaways from this analysis?",
                "How can this information be used in practice?",
                "What additional information would enhance this understanding?"
            ]


# ==================== FILE PARSER ====================
class FileParser:
    """Service for parsing various file formats"""
    
    def parse_dataset_file(self, file) -> Tuple[List[Dict[str, Any]], List[str]]:
        """Parse a dataset file and return records and schema"""
        filename = file.name.lower()
        
        if filename.endswith('.csv'):
            return self._parse_csv(file)
        elif filename.endswith('.json'):
            return self._parse_json(file)
        elif filename.endswith(('.xlsx', '.xls', '.ods')):
            return self._parse_excel(file)
        elif filename.endswith('.txt'):
            return self._parse_txt(file)
        else:
            raise ValueError(f"Unsupported file format: {filename}")
    
    def _parse_csv(self, file) -> Tuple[List[Dict[str, Any]], List[str]]:
        """Parse CSV file"""
        df = pd.read_csv(file)
        df.columns = df.columns.str.strip()
        records = df.to_dict('records')
        schema = list(df.columns)
        return records, schema
    
    def _parse_json(self, file) -> Tuple[List[Dict[str, Any]], List[str]]:
        """Parse JSON file"""
        content = file.read().decode('utf-8')
        data = json.loads(content)
        
        if isinstance(data, list):
            records = data
        elif isinstance(data, dict):
            for value in data.values():
                if isinstance(value, list):
                    records = value
                    break
            else:
                records = [data]
        else:
            raise ValueError("JSON must be an array or object")
        
        schema = list(records[0].keys()) if records else []
        return records, schema
    
    def _parse_excel(self, file) -> Tuple[List[Dict[str, Any]], List[str]]:
        """Parse Excel/ODS file"""
        df = pd.read_excel(file)
        df.columns = df.columns.str.strip()
        records = df.to_dict('records')
        schema = list(df.columns)
        return records, schema
    
    def _parse_txt(self, file) -> Tuple[List[Dict[str, Any]], List[str]]:
        """Parse TXT file as line-separated records"""
        content = file.read().decode('utf-8')
        lines = [line.strip() for line in content.split('\n') if line.strip()]
        records = [{'line': i + 1, 'content': line} for i, line in enumerate(lines)]
        schema = ['line', 'content']
        return records, schema
    
    def parse_template_file(self, file) -> str:
        """Parse a template file and return its content"""
        filename = file.name.lower()
        
        if filename.endswith(('.txt', '.md')):
            return file.read().decode('utf-8')
        elif filename.endswith('.docx'):
            return self._parse_docx(file)
        else:
            raise ValueError(f"Unsupported template format: {filename}")
    
    def _parse_docx(self, file) -> str:
        """Parse DOCX file using mammoth"""
        try:
            import mammoth
            result = mammoth.extract_raw_text(file)
            return result.value
        except ImportError:
            raise Exception("mammoth library not installed. Install with: pip install mammoth")


# ==================== TEMPLATE ENGINE ====================
class TemplateEngine:
    """Template rendering engine"""
    
    def render(self, template: str, data: Dict[str, Any]) -> str:
        """Render template with data"""
        result = template
        for key, value in data.items():
            placeholder = f"{{{{{key}}}}}"
            result = result.replace(placeholder, str(value))
        return result


# ==================== DOCUMENT GENERATOR ====================
class DocumentGenerator:
    """Document generation service"""
    
    def generate(self, template: str, dataset: List[Dict[str, Any]], output_format: str) -> List[Dict[str, str]]:
        """Generate documents from template and dataset"""
        engine = TemplateEngine()
        docs = []
        
        for idx, record in enumerate(dataset):
            content = engine.render(template, record)
            filename = self._generate_filename(record, idx, output_format)
            docs.append({
                'filename': filename,
                'content': content,
                'format': output_format
            })
        
        return docs
    
    def _generate_filename(self, record: Dict[str, Any], idx: int, format: str) -> str:
        """Generate filename from record"""
        # Try to use common fields for filename
        for field in ['name', 'title', 'id', 'filename']:
            if field in record:
                safe_name = str(record[field]).replace(' ', '_')[:50]
                return f"{safe_name}.{format}"
        return f"document_{idx + 1}.{format}"
    
    def export_all(self, docs: List[Dict[str, str]], output_format: str) -> bytes:
        """Export all documents as ZIP"""
        zip_buffer = io.BytesIO()
        
        with zipfile.ZipFile(zip_buffer, 'w', zipfile.ZIP_DEFLATED) as zip_file:
            for doc in docs:
                if output_format == 'docx':
                    content = self._convert_to_docx(doc['content'])
                else:
                    content = doc['content'].encode('utf-8')
                
                zip_file.writestr(doc['filename'], content)
        
        zip_buffer.seek(0)
        return zip_buffer.getvalue()
    
    def _convert_to_docx(self, content: str) -> bytes:
        """Convert text to DOCX format"""
        try:
            from docx import Document
            doc = Document()
            for paragraph in content.split('\n'):
                doc.add_paragraph(paragraph)
            
            buffer = io.BytesIO()
            doc.save(buffer)
            buffer.seek(0)
            return buffer.getvalue()
        except ImportError:
            # Fallback to plain text if python-docx not available
            return content.encode('utf-8')


# ==================== VISUALIZATION ====================
def create_pipeline_visualization(results: List[Dict], total_agents: int):
    """Create an interactive pipeline visualization"""
    fig = go.Figure()
    
    x_pos, y_pos, colors, text = [], [], [], []
    
    for idx, result in enumerate(results):
        x_pos.append(idx)
        y_pos.append(1)
        colors.append('#28A745' if result['status'] == 'success' else '#DC3545')
        text.append(f"{result['agent']}<br>{result['model']}")
    
    for idx in range(len(results), total_agents):
        x_pos.append(idx)
        y_pos.append(1)
        colors.append('#E6E6FA')
        text.append(f"Agent {idx + 1}<br>Pending")
    
    fig.add_trace(go.Scatter(
        x=x_pos,
        y=y_pos,
        mode='markers+text',
        marker=dict(size=50, color=colors, line=dict(color='#9370DB', width=3)),
        text=text,
        textposition="bottom center",
        textfont=dict(size=10),
        hoverinfo='text'
    ))
    
    for i in range(len(x_pos) - 1):
        fig.add_trace(go.Scatter(
            x=[x_pos[i], x_pos[i+1]],
            y=[y_pos[i], y_pos[i+1]],
            mode='lines',
            line=dict(color='#9370DB', width=2),
            hoverinfo='none',
            showlegend=False
        ))
    
    fig.update_layout(
        title="Agent Pipeline Flow",
        showlegend=False,
        height=200,
        margin=dict(l=20, r=20, t=40, b=20),
        xaxis=dict(showgrid=False, zeroline=False, showticklabels=False),
        yaxis=dict(showgrid=False, zeroline=False, showticklabels=False),
        plot_bgcolor='rgba(0,0,0,0)',
        paper_bgcolor='rgba(0,0,0,0)'
    )
    
    return fig


# ==================== SESSION STATE INITIALIZATION ====================
def init_session_state():
    if 'dataset' not in st.session_state:
        st.session_state.dataset = None
    if 'schema' not in st.session_state:
        st.session_state.schema = []
    if 'template' not in st.session_state:
        st.session_state.template = ""
    if 'generated_docs' not in st.session_state:
        st.session_state.generated_docs = []
    if 'agents_config' not in st.session_state:
        st.session_state.agents_config = load_default_agents()
    if 'pipeline_results' not in st.session_state:
        st.session_state.pipeline_results = []
    if 'follow_up_questions' not in st.session_state:
        st.session_state.follow_up_questions = []
    if 'output_format' not in st.session_state:
        st.session_state.output_format = 'txt'
    if 'ai_service' not in st.session_state:
        st.session_state.ai_service = None


def load_default_agents():
    """Load default agent configurations"""
    try:
        with open('agents.yaml', 'r') as f:
            return yaml.safe_load(f)
    except FileNotFoundError:
        return {
            'agents': [
                {
                    'name': 'Summarizer',
                    'model': 'gemini-2.5-flash',
                    'provider': 'gemini',
                    'system_prompt': 'You are an expert summarizer. Create concise, informative summaries.',
                    'user_prompt': 'Summarize the following text:\n\n{{input}}',
                    'temperature': 0.3,
                    'max_tokens': 1000,
                    'top_p': 0.9
                },
                {
                    'name': 'Style Rewriter',
                    'model': 'gpt-4o-mini',
                    'provider': 'openai',
                    'system_prompt': 'You are a professional writer. Rewrite content to be clear, engaging, and professional.',
                    'user_prompt': 'Rewrite the following text in a professional tone:\n\n{{input}}',
                    'temperature': 0.7,
                    'max_tokens': 1500,
                    'top_p': 0.95
                },
                {
                    'name': 'JSON Converter',
                    'model': 'grok-3-mini',
                    'provider': 'grok',
                    'system_prompt': 'You are a data structuring expert. Convert text into well-structured JSON.',
                    'user_prompt': 'Convert the following text into structured JSON:\n\n{{input}}',
                    'temperature': 0.2,
                    'max_tokens': 2000,
                    'top_p': 0.9
                }
            ]
        }


# ==================== MAIN APPLICATION ====================
init_session_state()

# Sidebar configuration
with st.sidebar:
    st.title("🌸 Flora Edition")
    st.divider()
    
    # API Configuration
    st.subheader("🔑 API Configuration")
    
    with st.expander("Configure API Keys", expanded=False):
        gemini_key = st.text_input("Gemini API Key", type="password", value=os.getenv("GEMINI_API_KEY", ""))
        openai_key = st.text_input("OpenAI API Key", type="password", value=os.getenv("OPENAI_API_KEY", ""))
        grok_key = st.text_input("Grok API Key (xAI)", type="password", value=os.getenv("XAI_API_KEY", ""))
        
        if st.button("💾 Save API Keys"):
            if gemini_key or openai_key or grok_key:
                st.session_state.ai_service = AIService(
                    gemini_key=gemini_key,
                    openai_key=openai_key,
                    grok_key=grok_key
                )
                st.success("✅ API keys configured!")
            else:
                st.warning("⚠️ Please provide at least one API key")
    
    st.divider()
    
    # Output format selection
    st.subheader("📄 Output Format")
    st.session_state.output_format = st.selectbox(
        "Document Format",
        options=['txt', 'markdown', 'docx'],
        index=0
    )
    
    st.divider()
    
    # Statistics
    if st.session_state.dataset is not None:
        st.subheader("📊 Statistics")
        col1, col2 = st.columns(2)
        with col1:
            st.metric("Records", len(st.session_state.dataset))
        with col2:
            st.metric("Columns", len(st.session_state.schema))
        
        if st.session_state.generated_docs:
            st.metric("Documents", len(st.session_state.generated_docs))
    
    st.divider()
    st.markdown("Made with 💜 by Flora Edition")
    st.markdown(f"Version 1.0.0")

# Main content
st.title("🌸 Agentic Docs Builder - Flora Edition")
st.markdown("*Automated document generation with AI-powered processing pipeline*")

# Create tabs
tab1, tab2, tab3, tab4, tab5 = st.tabs([
    "📁 Data Ingestion",
    "📝 Template Definition",
    "📄 Generated Documents",
    "🤖 Agent Configuration",
    "🚀 Pipeline Execution"
])

# TAB 1: Data Ingestion
with tab1:
    st.header("📁 Data Ingestion")
    st.markdown("Upload your structured dataset (CSV, JSON, XLSX, ODS, TXT)")
    
    uploaded_file = st.file_uploader(
        "Choose a dataset file",
        type=['csv', 'json', 'xlsx', 'ods', 'txt'],
        help="Upload a structured dataset file"
    )
    
    if uploaded_file is not None:
        try:
            with st.spinner("🔄 Parsing dataset..."):
                parser = FileParser()
                dataset, schema = parser.parse_dataset_file(uploaded_file)
                st.session_state.dataset = dataset
                st.session_state.schema = schema
            
            st.markdown('<div class="status-box status-success">✅ Dataset loaded successfully!</div>', unsafe_allow_html=True)
            
            # Display dataset preview
            st.subheader("📊 Dataset Preview (First 10 Records)")
            df = pd.DataFrame(dataset[:10])
            st.dataframe(df, use_container_width=True, height=400)
            
            # Dataset statistics
            col1, col2, col3 = st.columns(3)
            with col1:
                st.markdown('<div class="metric-card">', unsafe_allow_html=True)
                st.metric("Total Records", len(dataset))
                st.markdown('</div>', unsafe_allow_html=True)
            with col2:
                st.markdown('<div class="metric-card">', unsafe_allow_html=True)
                st.metric("Columns", len(schema))
                st.markdown('</div>', unsafe_allow_html=True)
            with col3:
                st.markdown('<div class="metric-card">', unsafe_allow_html=True)
                st.metric("Data Types", len(set(type(v).__name__ for row in dataset for v in row.values())))
                st.markdown('</div>', unsafe_allow_html=True)
            
            # Column information
            with st.expander("📋 Column Details"):
                col_info = []
                for col in schema:
                    sample_values = [row.get(col) for row in dataset[:5] if col in row]
                    col_info.append({
                        'Column': col,
                        'Type': type(sample_values[0]).__name__ if sample_values else 'Unknown',
                        'Sample': str(sample_values[0])[:50] if sample_values else 'N/A'
                    })
                st.table(pd.DataFrame(col_info))
            
        except Exception as e:
            st.markdown(f'<div class="status-box status-error">❌ Error: {str(e)}</div>', unsafe_allow_html=True)

# TAB 2: Template Definition
with tab2:
    st.header("📝 Template Definition")
    st.markdown("Define your document template using `{{column_name}}` placeholders")
    
    col1, col2 = st.columns([1, 1])
    
    with col1:
        st.subheader("Template Input")
        
        # Upload template file
        template_file = st.file_uploader(
            "Upload template file (optional)",
            type=['txt', 'md', 'docx'],
            key="template_upload"
        )
        
        if template_file is not None:
            try:
                parser = FileParser()
                template_content = parser.parse_template_file(template_file)
                st.session_state.template = template_content
                st.success("✅ Template loaded from file!")
            except Exception as e:
                st.error(f"❌ Error loading template: {str(e)}")
        
        # Template editor
        st.session_state.template = st.text_area(
            "Template Content",
            value=st.session_state.template,
            height=400,
            placeholder="Enter your template here...\n\nExample:\nName: {{name}}\nEmail: {{email}}\nDescription: {{description}}"
        )
        
        if st.session_state.schema:
            with st.expander("📋 Available Columns"):
                st.write(", ".join([f"`{{{{{col}}}}}`" for col in st.session_state.schema]))
    
    with col2:
        st.subheader("Live Preview")
        
        if st.session_state.template and st.session_state.dataset:
            try:
                engine = TemplateEngine()
                preview = engine.render(st.session_state.template, st.session_state.dataset[0])
                st.markdown("**Preview with first record:**")
                st.text_area("", value=preview, height=400, disabled=True, key="preview")
            except Exception as e:
                st.error(f"Preview error: {str(e)}")
        else:
            st.info("ℹ️ Upload data and create a template to see preview")
    
    st.divider()
    
    # Generate documents button
    if st.button("🚀 Generate Documents", type="primary", use_container_width=True):
        if not st.session_state.template:
            st.error("❌ Please create a template first")
        elif not st.session_state.dataset:
            st.error("❌ Please upload a dataset first")
        else:
            with st.spinner("🔄 Generating documents..."):
                try:
                    generator = DocumentGenerator()
                    docs = generator.generate(
                        st.session_state.template,
                        st.session_state.dataset,
                        st.session_state.output_format
                    )
                    st.session_state.generated_docs = docs
                    st.success(f"✅ Generated {len(docs)} documents!")
                    st.balloons()
                except Exception as e:
                    st.error(f"❌ Generation error: {str(e)}")

# TAB 3: Generated Documents
with tab3:
    st.header("📄 Generated Documents")
    
    if not st.session_state.generated_docs:
        st.info("ℹ️ No documents generated yet. Go to 'Template Definition' to create documents.")
    else:
        st.success(f"✅ {len(st.session_state.generated_docs)} documents ready")
        
        # Export all button
        col1, col2, col3 = st.columns([2, 1, 1])
        with col1:
            if st.button("📦 Export All Documents", type="primary"):
                generator = DocumentGenerator()
                zip_buffer = generator.export_all(st.session_state.generated_docs, st.session_state.output_format)
                st.download_button(
                    label="⬇️ Download ZIP",
                    data=zip_buffer,
                    file_name=f"documents_{datetime.now().strftime('%Y%m%d_%H%M%S')}.zip",
                    mime="application/zip"
                )
        
        st.divider()
        
        # Search and filter
        search = st.text_input("🔍 Search documents", placeholder="Search by filename or content...")
        
        # Display documents
        for idx, doc in enumerate(st.session_state.generated_docs):
            if search and search.lower() not in doc['filename'].lower() and search.lower() not in doc['content'].lower():
                continue
            
            with st.expander(f"📄 {doc['filename']}", expanded=False):
                edited_content = st.text_area(
                    "Content",
                    value=doc['content'],
                    height=200,
                    key=f"doc_{idx}"
                )
                st.session_state.generated_docs[idx]['content'] = edited_content
                
                col1, col2 = st.columns([1, 1])
                with col1:
                    st.download_button(
                        label="⬇️ Download",
                        data=edited_content,
                        file_name=doc['filename'],
